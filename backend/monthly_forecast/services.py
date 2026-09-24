"""
单月签排发填报 — 业务逻辑

数据来源：
- 本模块自己的 4 张 mf_ 表（读写）
- 既有 cc_ledger_contract / cc_country_mapping / cc_trade_module_data（只读，绝不修改）
"""
import json
import os
import tempfile
from datetime import date, datetime

from sqlalchemy import or_

from models import db
from monthly_forecast.models import MonthlyInput, ShipSelection, Submission, UserScope

# 既有表 — 只读引用
from dashboards.contract_completion.models import (
    CountryMapping, LedgerContract, TradeModuleData,
)

GAIZAO_MODULE = '改造'


# 月份工具

def default_month():
    """默认填报月份 = 下个月（YYYY-MM）"""
    today = date.today()
    year, month = today.year, today.month + 1
    if month > 12:
        year, month = year + 1, 1
    return f'{year}-{month:02d}'


def normalize_month(value):
    """校验并规整 YYYY-MM，非法返回 None"""
    text = (value or '').strip()
    if len(text) != 7 or text[4] != '-':
        return None
    try:
        year = int(text[:4])
        month = int(text[5:])
    except ValueError:
        return None
    if not (2000 <= year <= 2100 and 1 <= month <= 12):
        return None
    return f'{year}-{month:02d}'


def month_options():
    """可选月份：当月前后若干月 + 已有填报数据的月份"""
    today = date.today()
    months = set()
    for delta in range(-3, 4):
        month = today.month + delta
        year = today.year
        while month < 1:
            month += 12
            year -= 1
        while month > 12:
            month -= 12
            year += 1
        months.add(f'{year}-{month:02d}')
    for row in db.session.query(MonthlyInput.data_month).distinct().all():
        if row[0]:
            months.add(row[0])
    for row in db.session.query(ShipSelection.data_month).distinct().all():
        if row[0]:
            months.add(row[0])
    return sorted(months)


# 模块与归属判定

def load_module_map():
    """{国家: 模块名}（只读 cc_country_mapping）"""
    return {m.country: (m.module_name or '') for m in CountryMapping.query.all()}


def resolve_module(contract, module_map):
    """合同归属模块，三档优先级：改造 > 国家映射 > mapped_module"""
    if contract.product_type and GAIZAO_MODULE in str(contract.product_type):
        return GAIZAO_MODULE
    module = module_map.get(contract.country or '')
    if module:
        return module
    return contract.mapped_module or ''


def all_modules():
    """全部模块名：映射表模块 + 商贸模块 + 改造 + 已填报过的模块"""
    modules = {m.module_name for m in CountryMapping.query.all() if m.module_name}
    modules.update(t.module_name for t in TradeModuleData.query.all() if t.module_name)
    modules.add(GAIZAO_MODULE)
    modules.update(
        row[0] for row in db.session.query(MonthlyInput.module_name).distinct().all() if row[0]
    )
    return sorted(modules)


# 名单（scope）

def scope_modules(user_id):
    """当前用户可填报的模块列表"""
    rows = UserScope.query.filter_by(user_id=user_id).all()
    return [r.module_name for r in rows]


def has_scope(user_id):
    return UserScope.query.filter_by(user_id=user_id).count() > 0


# 候选池

def _date_str(value):
    return value.strftime('%Y-%m-%d') if value else ''


def _ship_within_month(value, data_month):
    """整梯发货日期（台账 AV 列）是否不晚于填报月份；None = 未确定，保留"""
    if value is None:
        return True
    return value.strftime('%Y-%m') <= data_month


def list_candidates(module_name, data_month=None, keyword=None):
    """发货候选池：组A日期为空 + 排产日期不为空 + 未作废 + 台账来源
    额外：整梯发货日期（台账 AV 列）晚于填报月份的梯号不进候选（空值保留）
    排序：排产日期倒序（日期近的在前），同日按合同号 / 梯号升序
    """
    contracts = LedgerContract.query.filter(
        LedgerContract.source == 'ledger',
        or_(LedgerContract.product_status.is_(None), LedgerContract.product_status != '已作废'),
        LedgerContract.delivery_date.is_(None),
        LedgerContract.schedule_date.isnot(None),
    ).all()

    module_map = load_module_map()
    rows = [c for c in contracts if resolve_module(c, module_map) == module_name]

    # 整梯发货日期（AV）不得晚于填报月份；与 delivery_date（AZ 组A日期）无关
    if data_month:
        rows = [c for c in rows if _ship_within_month(c.whole_ship_date, data_month)]

    # 稳定排序：先按合同号 / 梯号升序，再按排产日期倒序
    rows.sort(key=lambda c: (c.contract_no or '', c.ladder_no or ''))
    rows.sort(key=lambda c: c.schedule_date or date.min, reverse=True)

    if keyword:
        text = str(keyword).strip().lower()
        if text:
            rows = [c for c in rows if text in (
                (c.contract_no or '') + ' ' + (c.ladder_no or '') + ' ' + (c.project_name or '')
            ).lower()]

    return [{
        'contract_no': c.contract_no or '',
        'ladder_no': c.ladder_no or '',
        'project_name': c.project_name or '',
        'unit_count': c.unit_count or 0,
        'amount_wan': round((c.contract_amount_rmb or 0) / 10000.0, 2),
        'schedule_date': _date_str(c.schedule_date),
        'whole_ship_date': _date_str(c.whole_ship_date),
        'product_status': c.product_status or '',
    } for c in rows]


# 梯号校验

_FULLWIDTH_START = 0xFF01
_FULLWIDTH_END = 0xFF5E
_FULLWIDTH_OFFSET = 0xFEE0


def normalize_ladder(value):
    """梯号归一化：全角转半角、去掉空白、字母转大写

    台账里的写法固定是 合同编号/编号#（半角、无空格），但手工输入常带
    全角 ＃ ／ － 或空格，原样比对会一律判成「台账中无此梯号」。
    """
    chars = []
    for char in str(value or ''):
        code = ord(char)
        if code == 0x3000:
            continue
        if _FULLWIDTH_START <= code <= _FULLWIDTH_END:
            char = chr(code - _FULLWIDTH_OFFSET)
        if char.isspace():
            continue
        chars.append(char)
    return ''.join(chars).upper()


def find_ledger_ladder(ladder_no):
    """按梯号找台账行：先原样比对，再按归一化写法比对"""
    ladder = str(ladder_no or '').strip()
    if not ladder:
        return None
    row = LedgerContract.query.filter(LedgerContract.ladder_no == ladder).first()
    if row is not None:
        return row
    normalized = normalize_ladder(ladder)
    if not normalized or normalized == ladder:
        return None
    return LedgerContract.query.filter(LedgerContract.ladder_no == normalized).first()


def _ladder_samples(contract_no, limit=10):
    """同一个合同在台账里已有的梯号，用于提示（不参与校验）"""
    if not contract_no:
        return []
    rows = LedgerContract.query.filter(
        LedgerContract.contract_no == contract_no
    ).all()
    return sorted({r.ladder_no for r in rows if r.ladder_no})[:limit]


def validate_ladder(module_name, data_month, ladder_no):
    """手工添加梯号校验，返回几种结果之一

    梯号必须能在台账里找到（可为全角 / 带空格 / 小写写法，会先归一化）；
    匹配成功后一律回写台账里的标准写法，避免存进明细后又对不上。
    """
    ladder = str(ladder_no or '').strip()
    if not ladder:
        return {'ok': False, 'code': 'EMPTY', 'msg': '请输入梯号', 'row': None, 'suggestions': []}

    contract = find_ledger_ladder(ladder)
    if contract is None:
        samples = _ladder_samples(ladder.split('/')[0].strip())
        if samples:
            return {
                'ok': False, 'code': 'NEED_LADDER', 'row': None, 'suggestions': samples,
                'msg': '请填写完整梯号，例如 ' + '、'.join(samples[:3]),
            }
        return {
            'ok': False, 'code': 'NOT_FOUND', 'row': None, 'suggestions': [],
            'msg': '台账中无此梯号（请核对标的编号写法，如 DHT-261102T/1#）',
        }

    owner = resolve_module(contract, load_module_map())
    if owner != module_name:
        return {'ok': False, 'code': 'OTHER_MODULE',
                'msg': '不属于本模块（台账归属：' + (owner or '未知') + '）',
                'row': None, 'suggestions': []}

    exists = ShipSelection.query.filter_by(
        data_month=data_month, module_name=module_name, ladder_no=contract.ladder_no
    ).first()
    if exists:
        return {'ok': False, 'code': 'DUPLICATE', 'msg': '该梯号已在列表中',
                'row': None, 'suggestions': []}

    return {'ok': True, 'code': 'OK', 'msg': '已添加', 'row': {
        'contract_no': contract.contract_no or '',
        'ladder_no': contract.ladder_no or ladder,
        'project_name': contract.project_name or '',
        'unit_count': contract.unit_count or 0,
        'amount_wan': round((contract.contract_amount_rmb or 0) / 10000.0, 2),
        'schedule_date': _date_str(contract.schedule_date),
        'is_manual': True,
    }, 'suggestions': []}


# 填报读写

def _enrich_ship_rows(ships):
    """给勾选明细补上台账侧的项目名称 / 排产日期（只读查询）"""
    ladders = [s.ladder_no for s in ships if s.ladder_no]
    contracts = {}
    if ladders:
        for c in LedgerContract.query.filter(LedgerContract.ladder_no.in_(ladders)).all():
            contracts[c.ladder_no] = c
    result = []
    for s in ships:
        row = s.to_dict()
        contract = contracts.get(s.ladder_no)
        row['project_name'] = (contract.project_name or '') if contract else ''
        row['schedule_date'] = _date_str(contract.schedule_date) if contract else ''
        result.append(row)
    return result


def _sum_ships(ships):
    """发货合计：台数 / 金额（万元）"""
    units = round(sum(s.unit_count or 0 for s in ships), 2)
    amount = round(sum(s.amount_rmb or 0 for s in ships) / 10000.0, 2)
    return units, amount


def _ship_totals(data_month, module_name):
    """发货合计 = 梯号明细合计 + 无梯号手工预估"""
    ships = ShipSelection.query.filter_by(
        data_month=data_month, module_name=module_name
    ).all()
    units, amount = _sum_ships(ships)
    record = MonthlyInput.query.filter_by(
        data_month=data_month, module_name=module_name
    ).first()
    manual_units = (record.ship_manual_units if record else 0.0) or 0.0
    manual_amount = (record.ship_manual_amount if record else 0.0) or 0.0
    return round(units + manual_units, 2), round(amount + manual_amount, 2)


def get_entry(data_month, module_name):
    """读取某月某模块的填报内容"""
    record = MonthlyInput.query.filter_by(
        data_month=data_month, module_name=module_name
    ).first()
    ships = ShipSelection.query.filter_by(
        data_month=data_month, module_name=module_name
    ).order_by(ShipSelection.id).all()
    ship_units, ship_amount = _ship_totals(data_month, module_name)

    base = record.to_dict() if record else {
        'data_month': data_month, 'module_name': module_name,
        'sign_units': 0.0, 'sign_amount': 0.0, 'prod_units': 0.0, 'prod_amount': 0.0,
        'ship_manual_units': 0.0, 'ship_manual_amount': 0.0,
        'status': 'draft', 'submitted_at': '', 'updated_by_name': '', 'updated_at': '',
    }
    base['ship_units'] = ship_units
    base['ship_amount'] = ship_amount
    base['ship_selections'] = _enrich_ship_rows(ships)
    return base


def _to_float(value):
    if value in (None, ''):
        return 0.0
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return 0.0


def save_draft(data_month, module_name, payload, user):
    """保存草稿：upsert 签排产 + 全量替换发货勾选明细

    发货明细只接受梯号，台数与金额由服务端从台账重新快照，不信任前端数值；
    无梯号发货预估（ship_manual_units / ship_manual_amount）为纯手填，直接采用。
    不属于本模块或台账中不存在的梯号会被拒绝并返回 rejected 列表。
    """
    payload = payload or {}

    record = MonthlyInput.query.filter_by(
        data_month=data_month, module_name=module_name
    ).first()
    if record is None:
        record = MonthlyInput(data_month=data_month, module_name=module_name)
        db.session.add(record)

    record.sign_units = _to_float(payload.get('sign_units'))
    record.sign_amount = _to_float(payload.get('sign_amount'))
    record.prod_units = _to_float(payload.get('prod_units'))
    record.prod_amount = _to_float(payload.get('prod_amount'))
    record.ship_manual_units = _to_float(payload.get('ship_manual_units'))
    record.ship_manual_amount = _to_float(payload.get('ship_manual_amount'))
    record.updated_by = getattr(user, 'id', None)
    record.updated_by_name = getattr(user, 'real_name', '') or getattr(user, 'username', '') or ''
    record.updated_at = datetime.now()

    ShipSelection.query.filter_by(
        data_month=data_month, module_name=module_name
    ).delete(synchronize_session=False)

    selections = payload.get('ship_selections') or []
    ladders = []
    manual_flags = {}
    for item in selections:
        item = item or {}
        ladder = str(item.get('ladder_no') or '').strip()
        if not ladder or ladder in manual_flags:
            continue
        ladders.append(ladder)
        manual_flags[ladder] = bool(item.get('is_manual'))

    rejected = []
    if ladders:
        contracts = {}
        for c in LedgerContract.query.filter(LedgerContract.ladder_no.in_(ladders)).all():
            contracts[c.ladder_no] = c
        # 手工输入的 全角 / 空格 写法，按归一化再找一次
        pending = {}
        for ladder in ladders:
            if ladder in contracts:
                continue
            normalized = normalize_ladder(ladder)
            if normalized and normalized not in pending:
                pending[normalized] = ladder
        if pending:
            for c in LedgerContract.query.filter(
                LedgerContract.ladder_no.in_(list(pending))
            ).all():
                contracts[pending[c.ladder_no]] = c
        module_map = load_module_map()
        for ladder in ladders:
            contract = contracts.get(ladder)
            if contract is None:
                rejected.append({'ladder_no': ladder, 'reason': '台账中无此梯号'})
                continue
            if resolve_module(contract, module_map) != module_name:
                rejected.append({'ladder_no': ladder, 'reason': '不属于本模块'})
                continue
            db.session.add(ShipSelection(
                data_month=data_month,
                module_name=module_name,
                contract_no=contract.contract_no or '',
                ladder_no=contract.ladder_no or ladder,
                unit_count=contract.unit_count or 0,
                amount_rmb=contract.contract_amount_rmb or 0,
                is_manual=manual_flags[ladder],
                created_by=getattr(user, 'id', None),
                created_by_name=record.updated_by_name,
            ))

    db.session.commit()
    entry = get_entry(data_month, module_name)
    entry['rejected'] = rejected
    return entry


def submit(data_month, module_name, payload, user):
    """提交：先落库当前填报，再生成留存快照（版本号自增，历史不可覆盖）"""
    entry = save_draft(data_month, module_name, payload, user)

    record = MonthlyInput.query.filter_by(
        data_month=data_month, module_name=module_name
    ).first()
    record.status = 'submitted'
    record.submitted_at = datetime.now()

    last = Submission.query.filter_by(
        data_month=data_month, module_name=module_name
    ).order_by(Submission.version.desc()).first()
    version = (last.version + 1) if last else 1

    submission = Submission(
        data_month=data_month,
        module_name=module_name,
        version=version,
        payload=json.dumps({
            'sign_units': entry['sign_units'], 'sign_amount': entry['sign_amount'],
            'prod_units': entry['prod_units'], 'prod_amount': entry['prod_amount'],
            'ship_units': entry['ship_units'], 'ship_amount': entry['ship_amount'],
            'ship_manual_units': entry.get('ship_manual_units', 0.0),
            'ship_manual_amount': entry.get('ship_manual_amount', 0.0),
            'ship_selections': entry['ship_selections'],
        }, ensure_ascii=False),
        sign_units=entry['sign_units'], sign_amount=entry['sign_amount'],
        prod_units=entry['prod_units'], prod_amount=entry['prod_amount'],
        ship_units=entry['ship_units'], ship_amount=entry['ship_amount'],
        submitted_by=getattr(user, 'id', None),
        submitted_by_name=record.updated_by_name,
    )
    db.session.add(submission)
    db.session.commit()

    entry['status'] = 'submitted'
    entry['version'] = version
    return entry


def list_submissions(data_month, module_name):
    """历史提交版本列表（不含 payload 明细）"""
    rows = Submission.query.filter_by(
        data_month=data_month, module_name=module_name
    ).order_by(Submission.version.desc()).all()
    return [r.to_dict() for r in rows]


# 汇总

METRIC_KEYS = ('sign_units', 'sign_amount', 'prod_units', 'prod_amount', 'ship_units', 'ship_amount')


def _module_row(data_month, module_name):
    """单个模块的签排发 6 列"""
    record = MonthlyInput.query.filter_by(
        data_month=data_month, module_name=module_name
    ).first()
    ships = ShipSelection.query.filter_by(
        data_month=data_month, module_name=module_name
    ).all()
    ship_units, ship_amount = _sum_ships(ships)
    ship_units = round(ship_units + ((record.ship_manual_units if record else 0.0) or 0.0), 2)
    ship_amount = round(ship_amount + ((record.ship_manual_amount if record else 0.0) or 0.0), 2)
    return {
        'module_name': module_name,
        'sign_units': (record.sign_units if record else 0.0) or 0.0,
        'sign_amount': (record.sign_amount if record else 0.0) or 0.0,
        'prod_units': (record.prod_units if record else 0.0) or 0.0,
        'prod_amount': (record.prod_amount if record else 0.0) or 0.0,
        'ship_units': ship_units,
        'ship_amount': ship_amount,
        'status': (record.status if record else '') or '',
        'updated_by_name': (record.updated_by_name if record else '') or '',
        'updated_at': record.updated_at.strftime('%Y-%m-%d %H:%M:%S') if record and record.updated_at else '',
    }


def get_summary(data_month, modules=None):
    """汇总：全部（或指定）模块的签排发 6 列 + 合计行"""
    module_list = list(modules) if modules else all_modules()
    rows = [_module_row(data_month, m) for m in module_list]

    total = {key: 0.0 for key in METRIC_KEYS}
    for row in rows:
        for key in METRIC_KEYS:
            total[key] = round(total[key] + (row.get(key) or 0.0), 2)
    total['module_name'] = '合计'
    total['is_total'] = True

    return {'data_month': data_month, 'rows': rows, 'total': total}


# 导出

FORECAST_HEADERS = ['模块', '签单台数', '签单金额(万元)', '排产台数', '排产金额(万元)', '发货台数', '发货金额(万元)']
SHIPPING_HEADERS = ['模块', '合同号', '梯号', '项目名称', '台数', '金额(万元)', '来源', '填报人']


def _new_workbook():
    from openpyxl import Workbook
    return Workbook()


def _style_header(ws, count):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    font = Font(name='微软雅黑', bold=True, size=10)
    fill = PatternFill('solid', fgColor='EEF2F7')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
    for col in range(1, count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = font
        cell.fill = fill
        cell.alignment = center
        cell.border = border
    ws.freeze_panes = 'A2'


def _set_widths(ws, widths):
    from openpyxl.utils import get_column_letter
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def export_forecast(data_month, modules=None):
    """Excel 1：预计签排发台数金额"""
    summary = get_summary(data_month, modules=modules)
    wb = _new_workbook()
    ws = wb.active
    ws.title = '签排发预计'
    ws.append(FORECAST_HEADERS)

    for row in summary['rows']:
        ws.append([
            row['module_name'],
            row['sign_units'], row['sign_amount'],
            row['prod_units'], row['prod_amount'],
            row['ship_units'], row['ship_amount'],
        ])

    total = summary['total']
    ws.append([
        '合计',
        total['sign_units'], total['sign_amount'],
        total['prod_units'], total['prod_amount'],
        total['ship_units'], total['ship_amount'],
    ])

    _style_header(ws, len(FORECAST_HEADERS))
    _set_widths(ws, [24, 12, 16, 12, 16, 12, 16])

    path = os.path.join(tempfile.gettempdir(), 'mf_forecast_' + data_month + '.xlsx')
    wb.save(path)
    return path


def export_shipping(data_month, modules=None):
    """Excel 2：预计能发货的合同号梯号明细"""
    module_list = list(modules) if modules else all_modules()
    ships = ShipSelection.query.filter(
        ShipSelection.data_month == data_month,
        ShipSelection.module_name.in_(module_list),
    ).order_by(ShipSelection.module_name, ShipSelection.ladder_no).all()

    ladders = [s.ladder_no for s in ships if s.ladder_no]
    contracts = {}
    if ladders:
        for c in LedgerContract.query.filter(LedgerContract.ladder_no.in_(ladders)).all():
            contracts[c.ladder_no] = c

    wb = _new_workbook()
    ws = wb.active
    ws.title = '预计发货明细'
    ws.append(SHIPPING_HEADERS)

    for s in ships:
        contract = contracts.get(s.ladder_no)
        ws.append([
            s.module_name,
            s.contract_no or '',
            s.ladder_no or '',
            (contract.project_name or '') if contract else '',
            s.unit_count or 0,
            round((s.amount_rmb or 0) / 10000.0, 2),
            '手工输入' if s.is_manual else '候选勾选',
            s.created_by_name or '',
        ])

    _style_header(ws, len(SHIPPING_HEADERS))
    _set_widths(ws, [24, 18, 22, 32, 8, 14, 12, 12])

    path = os.path.join(tempfile.gettempdir(), 'mf_shipping_' + data_month + '.xlsx')
    wb.save(path)
    return path
