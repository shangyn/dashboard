import time, glob, sys
path = glob.glob(r"E:\System\backend\uploads\contract_completion\contract_ledger\*.xlsx")[0]
import openpyxl
t0 = time.time()
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb.active
print("active:", ws.title, "declared dims:", ws.calculate_dimension(), "max_row:", ws.max_row, "max_col:", ws.max_column)
sys.stdout.flush()
hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
names = [str(v).strip() for v in hdr]
print("header cells:", len(names), "non-empty:", sum(1 for n in names if n))
print("first 25:", [n for n in names if n][:25])
sys.stdout.flush()
need = ["合同编号", "合同号", "签订日期", "排产日期", "国家", "业务员"]
found = {k: (k in names) for k in need}
print("key columns present:", found)
t1 = time.time()
n = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    n += 1
t2 = time.time()
print(f"load {t1-t0:.1f}s  iterate {t2-t1:.1f}s  rows={n}")
