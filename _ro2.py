import time, glob, sys
path = glob.glob(r"E:\System\backend\uploads\contract_completion\contract_ledger\*.xlsx")[0]
import openpyxl
t0 = time.time()
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb.active
ws.reset_dimensions()
print("after reset_dimensions -> max_row:", ws.max_row, "max_col:", ws.max_column)
sys.stdout.flush()
it = ws.iter_rows(values_only=True)
hdr = next(it)
names = [str(v).strip() if v is not None else "" for v in hdr]
print("header cells:", len(names), "non-empty:", sum(1 for n in names if n))
print("keys found:", [k for k in ["合同编号","合同号","签订日期","排产日期","国家","业务员","大区域","小区域","标的编号","梯号"] if k in names])
sys.stdout.flush()
n = 0
for row in it:
    n += 1
t1 = time.time()
print(f"stream all rows: {t1-t0:.1f}s  data rows={n}  max_col={ws.max_column}")
