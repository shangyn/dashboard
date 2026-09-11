import glob, sys
path = glob.glob(r"E:\System\backend\uploads\contract_completion\contract_ledger\*.xlsx")[0]
import openpyxl
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb.active
ws.reset_dimensions()
hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
hmap = {str(v).strip(): i+1 for i, v in enumerate(hdr) if v is not None and str(v).strip()}
ncols = len(hdr)
print("ncols(header):", ncols)
cols = ["合同编号", "标的编号", "签订日期", "排产日期", "组A日期", "国家", "业务员", "大区域", "小区域", "项目名称", "产品型号"]
for j, raw in enumerate(ws.iter_rows(min_row=2, max_col=ncols, values_only=True)):
    row = raw if len(raw) >= ncols else tuple(raw) + (None,) * (ncols - len(raw))
    out = {}
    for c in cols:
        idx = hmap.get(c)
        out[c] = row[idx-1] if idx else None
    print(len(row), out)
    if j >= 2:
        break
wb.close()
