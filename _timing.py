import time, os, sys, glob, ctypes, ctypes.wintypes as wt

path = glob.glob(r"E:\System\backend\uploads\contract_completion\contract_ledger\*.xlsx")[0]
print("file:", os.path.basename(path), round(os.path.getsize(path)/1024/1024, 1), "MB")
sys.stdout.flush()

class PMC(ctypes.Structure):
    _fields_ = [("cb", wt.DWORD), ("PageFaultCount", wt.DWORD),
                ("PeakWorkingSetSize", ctypes.c_size_t), ("WorkingSetSize", ctypes.c_size_t),
                ("QuotaPeakPagedPoolUsage", ctypes.c_size_t), ("QuotaPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t), ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                ("PagefileUsage", ctypes.c_size_t), ("PeakPagefileUsage", ctypes.c_size_t)]

def peak_mb():
    p = PMC(); p.cb = ctypes.sizeof(PMC)
    ctypes.windll.psapi.GetProcessMemoryInfo(ctypes.windll.kernel32.GetCurrentProcess(), ctypes.byref(p), p.cb)
    return round(p.PeakWorkingSetSize/1024/1024)

import openpyxl
t0 = time.time()
wb = openpyxl.load_workbook(path, data_only=True)
t1 = time.time()
print(f"load_workbook: {t1-t0:.1f}s  peak RAM: {peak_mb()} MB")
sys.stdout.flush()

ws = wb.active
print("sheet:", ws.title, "max_row:", ws.max_row, "max_col:", ws.max_column)
n = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    n += 1
t2 = time.time()
print(f"iter_rows: {t2-t1:.1f}s  rows={n}  peak RAM: {peak_mb()} MB")
print(f"TOTAL parse(read-only work): {t2-t0:.1f}s")
