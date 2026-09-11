# 复刻 parse_ledger_contracts 的解析逻辑（不写库），验证 read_only 版结果与生产记录一致
import glob, time
from datetime import datetime, date
import openpyxl

path = glob.glob(r"E:\System\backend\uploads\contract_completion\contract_ledger\*.xlsx")[0]

def _safe_str(v):
    return '' if v is None else str(v).strip()

def _safe_date_openpyxl(value):
    if value is None: return None
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    s = str(value).strip()
    if not s or s == '-': return None
    for fmt in ['%Y-%m-%d','%Y/%m/%d','%Y.%m.%d','%m/%d/%Y','%d/%m/%Y']:
        try: return datetime.strptime(s, fmt).date()
        except ValueError: continue
    return None

t0 = time.time()
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb.active
ws.reset_dimensions()
header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
header_map = {}
for c, v in enumerate(header_row, start=1):
    v = _safe_str(v)
    if v: header_map[v] = c
ncols = len(header_row)

def _h(*names):
    for n in names:
        col = header_map.get(n)
        if col is not None: return col
    return None

col_contract_no = _h('合同编号','合同号')
col_sign_date = header_map.get('签订日期')
col_schedule_date = header_map.get('排产日期')
col_delivery_date = _h('组A日期','实际发货日期','实际发运日期')
print("cols:", col_contract_no, col_sign_date, col_schedule_date, col_delivery_date)

CUTOFF_DATE = date(2013, 1, 1)
total = skipped_old = rows_seen = 0
for raw_row in ws.iter_rows(min_row=2, max_col=ncols, values_only=True):
    row = raw_row if len(raw_row) >= ncols else tuple(raw_row) + (None,) * (ncols - len(raw_row))
    rows_seen += 1
    sd = _safe_date_openpyxl(row[col_sign_date-1]) if col_sign_date else None
    scd = _safe_date_openpyxl(row[col_schedule_date-1]) if col_schedule_date else None
    dd = _safe_date_openpyxl(row[col_delivery_date-1]) if col_delivery_date else None
    if not ((sd and sd >= CUTOFF_DATE) or (scd and scd >= CUTOFF_DATE) or (dd and dd >= CUTOFF_DATE)):
        skipped_old += 1
        continue
    cn = _safe_str(row[col_contract_no-1]) if col_contract_no else ''
    if not cn: continue
    total += 1
print(f"rows_seen={rows_seen} imported={total} skipped_old={skipped_old} elapsed={time.time()-t0:.1f}s")
