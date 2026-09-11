import time, os, glob, sys
path = glob.glob(r"E:\System\backend\uploads\contract_completion\contract_ledger\*.xlsx")[0]
import openpyxl
t0 = time.time()
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active
cells = ws.max_row * ws.max_column
print(f"loaded in {time.time()-t0:.1f}s cells={cells}")
sys.stdout.flush()
import ctypes, ctypes.wintypes as wt
class PMC(ctypes.Structure):
    _fields_=[("cb",wt.DWORD),("PageFaultCount",wt.DWORD),("PeakWorkingSetSize",ctypes.c_size_t),
              ("WorkingSetSize",ctypes.c_size_t),("QuotaPeakPagedPoolUsage",ctypes.c_size_t),
              ("QuotaPagedPoolUsage",ctypes.c_size_t),("QuotaPeakNonPagedPoolUsage",ctypes.c_size_t),
              ("QuotaNonPagedPoolUsage",ctypes.c_size_t),("PagefileUsage",ctypes.c_size_t),
              ("PeakPagefileUsage",ctypes.c_size_t)]
p=PMC(); p.cb=ctypes.sizeof(PMC)
ok=ctypes.windll.psapi.GetProcessMemoryInfo(ctypes.windll.kernel32.GetCurrentProcess(), ctypes.byref(p), p.cb)
print("psapi ok=",ok,"peak_ws_MB=",round(p.PeakWorkingSetSize/1024/1024),"cur_ws_MB=",round(p.WorkingSetSize/1024/1024))
time.sleep(25)
